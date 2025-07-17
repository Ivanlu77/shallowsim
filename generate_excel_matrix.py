#!/usr/bin/env python3
"""
生成Excel格式的完整性能矩阵表格
包含所有TP×DP组合与序列长度组合的不同GPU卡的结果
"""

import shallowsim as sb
import pandas as pd
import itertools
import numpy as np
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows

def main():
    print("=" * 80)
    print("ShallowSim Prefill 完整性能矩阵 - Excel格式输出")
    print("=" * 80)
    
    # 1. 初始化参数
    args = sb.ModelArgs()
    gpu_dict = sb.get_gpu_info('./device/gpu_info.csv', print_console=False)
    
    # 测试参数配置
    seq_len_list = [4383, 8766, 17532, 35064, 70128]  # 序列长度列表
    kv_cache_rate = 0.563                              # KV缓存命中率
    tp_list = [1, 2, 4]                               # Tensor Parallel 配置
    dp_list = [8, 16, 32, 64, 128]                    # Data Parallel 配置
    
    print(f"GPU类型({len(gpu_dict)}): {list(gpu_dict.keys())}")
    print(f"序列长度({len(seq_len_list)}): {seq_len_list}")
    print(f"TP配置({len(tp_list)}): {tp_list}")
    print(f"DP配置({len(dp_list)}): {dp_list}")
    print(f"总组合数: {len(seq_len_list) * len(tp_list) * len(dp_list)} = {len(seq_len_list)}×{len(tp_list)}×{len(dp_list)}")
    print("=" * 80)
    
   
    results = []
    total_combinations = len(seq_len_list) * len(tp_list) * len(dp_list)
    current_count = 0
    
    for seq_len, tp, dp in itertools.product(seq_len_list, tp_list, dp_list):
        current_count += 1
        combo_name = f"Seq{seq_len}_TP{tp}×DP{dp}"
        print(f"[{current_count:2d}/{total_combinations}] {combo_name}")
        
        try:
            # 计算prefill时间
            detail, summary = sb.prefill_time(args, gpu_dict, seq_len, kv_cache_rate, tp=tp, dp=dp, print_console=False)
            
            # 计算吞吐量
            throughput_per_device = summary.apply(lambda x: seq_len/tp * (1000 / x)).loc['Sum']
            
            
            for gpu_name in gpu_dict.keys():
                if gpu_name in summary.columns:
                    results.append({
                        'GPU': gpu_name,
                        'seq_len': seq_len,
                        'tp': tp,
                        'dp': dp,
                        'combo_name': combo_name,
                        'tp_dp_combo': f"TP{tp}×DP{dp}",
                        'device_count': tp * dp,
                        'throughput': throughput_per_device[gpu_name],
                        'prefill_time_ms': summary.loc['Sum', gpu_name]
                    })
        
        except Exception as e:
            print(f" 错误: {e}")
            
            for gpu_name in gpu_dict.keys():
                results.append({
                    'GPU': gpu_name,
                    'seq_len': seq_len,
                    'tp': tp,
                    'dp': dp,
                    'combo_name': combo_name,
                    'tp_dp_combo': f"TP{tp}×DP{dp}",
                    'device_count': tp * dp,
                    'throughput': np.nan,
                    'prefill_time_ms': np.nan
                })
            continue
    
    # 3. 转换为DataFrame
    results_df = pd.DataFrame(results)
    print(f"\n完成测试，共收集 {len(results_df)} 条结果")
    
    # 4. 生成各种表格
    print("\n生成Excel表格...")
    
    # 4.1 完整性能矩阵
    performance_matrix = results_df.pivot_table(
        index='GPU',
        columns='combo_name',
        values='throughput',
        aggfunc='first'
    ).round(1)
    
    # 按组合名称排序列
    def sort_key(combo_name):
        parts = combo_name.split('_')
        seq_len = int(parts[0][3:])
        tp_dp = parts[1]
        tp = int(tp_dp.split('×')[0][2:])
        dp = int(tp_dp.split('×')[1][2:])
        return (seq_len, tp, dp)
    
    sorted_columns = sorted(performance_matrix.columns, key=sort_key)
    performance_matrix = performance_matrix[sorted_columns]
    
    # 4.2 按序列长度分组的表格
    seq_matrices = {}
    for seq_len in seq_len_list:
        seq_data = results_df[results_df['seq_len'] == seq_len]
        seq_matrix = seq_data.pivot_table(
            index='GPU',
            columns='tp_dp_combo',
            values='throughput',
            aggfunc='first'
        ).round(1)
        
        # 按TP×DP组合排序
        tp_dp_order = []
        for tp in tp_list:
            for dp in dp_list:
                combo = f"TP{tp}×DP{dp}"
                if combo in seq_matrix.columns:
                    tp_dp_order.append(combo)
        
        seq_matrix = seq_matrix[tp_dp_order]
        seq_matrices[seq_len] = seq_matrix
    
    # 4.3 统计分析表格
    best_configs = results_df.loc[results_df.groupby(['GPU', 'seq_len'])['throughput'].idxmax()]
    best_performance_table = best_configs.pivot_table(
        index='GPU',
        columns='seq_len',
        values='throughput',
        aggfunc='first'
    ).round(1)
    
    best_config_table = best_configs.pivot_table(
        index='GPU',
        columns='seq_len',
        values='tp_dp_combo',
        aggfunc='first'
    )
    
    best_device_table = best_configs.pivot_table(
        index='GPU',
        columns='seq_len',
        values='device_count',
        aggfunc='first'
    )
    
    # 4.4 综合统计
    gpu_best_overall = results_df.loc[results_df.groupby('GPU')['throughput'].idxmax()]
    seq_best_overall = results_df.loc[results_df.groupby('seq_len')['throughput'].idxmax()]
    combo_analysis = results_df.groupby('tp_dp_combo')['throughput'].agg(['mean', 'std', 'min', 'max']).round(1)
    
    # 5. 生成Excel文件
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    excel_filename = f"prefill_complete_analysis_{timestamp}.xlsx"
    
    print(f"正在生成Excel文件: {excel_filename}")
    
    with pd.ExcelWriter(excel_filename, engine='openpyxl') as writer:
        # 工作表1: 完整性能矩阵
        performance_matrix.to_excel(writer, sheet_name='完整性能矩阵', startrow=1)
        ws1 = writer.sheets['完整性能矩阵']
        ws1['A1'] = '完整性能矩阵 (tokens/sec/device)'
        ws1['A1'].font = Font(bold=True, size=14)
        
        # 工作表2-6: 按序列长度分组
        for seq_len in seq_len_list:
            sheet_name = f'Seq{seq_len}'
            seq_matrices[seq_len].to_excel(writer, sheet_name=sheet_name, startrow=1)
            ws = writer.sheets[sheet_name]
            ws['A1'] = f'序列长度 {seq_len} 的性能矩阵 (tokens/sec/device)'
            ws['A1'].font = Font(bold=True, size=14)
            
            # 添加设备数量信息
            row_start = seq_matrices[seq_len].shape[0] + 4
            ws[f'A{row_start}'] = '设备总数对照:'
            ws[f'A{row_start}'].font = Font(bold=True)
            
            col_idx = 2
            for tp in tp_list:
                for dp in dp_list:
                    device_count = tp * dp
                    ws.cell(row=row_start+1, column=col_idx, value=f"TP{tp}×DP{dp}")
                    ws.cell(row=row_start+2, column=col_idx, value=f"{device_count}设备")
                    col_idx += 1
        
        # 工作表7: 最佳性能汇总
        best_performance_table.to_excel(writer, sheet_name='最佳性能汇总', startrow=1)
        ws7 = writer.sheets['最佳性能汇总']
        ws7['A1'] = '各GPU最佳性能 (tokens/sec/device)'
        ws7['A1'].font = Font(bold=True, size=14)
        
        # 添加最佳配置信息
        start_row = best_performance_table.shape[0] + 4
        ws7[f'A{start_row}'] = '对应的最佳TP×DP配置:'
        ws7[f'A{start_row}'].font = Font(bold=True)
        
        for i, (idx, row) in enumerate(best_config_table.iterrows()):
            ws7.cell(row=start_row+1+i, column=1, value=idx)
            for j, val in enumerate(row):
                ws7.cell(row=start_row+1+i, column=2+j, value=val)
        
        # 工作表8: TP×DP组合统计
        combo_analysis.to_excel(writer, sheet_name='TPDP组合统计')
        ws8 = writer.sheets['TPDP组合统计']
        ws8['A1'] = 'TP×DP组合性能统计'
        ws8['A1'].font = Font(bold=True, size=14)
        
        # 工作表9: GPU最佳配置详情
        gpu_best_summary = gpu_best_overall[['GPU', 'combo_name', 'seq_len', 'tp', 'dp', 'device_count', 'throughput', 'prefill_time_ms']].round(1)
        gpu_best_summary = gpu_best_summary.sort_values('throughput', ascending=False)
        gpu_best_summary.to_excel(writer, sheet_name='GPU最佳配置', index=False)
        ws9 = writer.sheets['GPU最佳配置']
        ws9['A1'] = 'GPU最佳性能配置详情'
        ws9['A1'].font = Font(bold=True, size=14)
        
        # 工作表10: 原始数据
        results_df.to_excel(writer, sheet_name='原始数据', index=False)
        ws10 = writer.sheets['原始数据']
        ws10['A1'] = '完整原始测试数据'
        ws10['A1'].font = Font(bold=True, size=14)
        
        # 工作表11: 测试配置信息
        config_info = pd.DataFrame({
            '配置项': ['GPU类型数量', '序列长度列表', 'TP配置', 'DP配置', 'KV缓存命中率', '总测试组合数'],
            '值': [
                len(gpu_dict),
                str(seq_len_list),
                str(tp_list),
                str(dp_list),
                kv_cache_rate,
                len(seq_len_list) * len(tp_list) * len(dp_list)
            ]
        })
        config_info.to_excel(writer, sheet_name='测试配置', index=False)
        ws11 = writer.sheets['测试配置']
        ws11['A1'] = '测试配置信息'
        ws11['A1'].font = Font(bold=True, size=14)
        
        # 添加GPU列表
        gpu_list_df = pd.DataFrame({'GPU类型': list(gpu_dict.keys())})
        gpu_list_df.to_excel(writer, sheet_name='测试配置', startcol=4, index=False)
    
    print(f"✅ Excel文件生成完成: {excel_filename}")
    
    # 6. 显示简要统计
    print(f"\n📊 简要统计:")
    print(f"- 测试了 {len(gpu_dict)} 种GPU")
    print(f"- 每种GPU测试了 {len(seq_len_list) * len(tp_list) * len(dp_list)} 种配置组合")
    print(f"- 总共 {len(results_df)} 条测试记录")
    print(f"- Excel文件包含 11 个工作表")
    
    print(f"\n📋 Excel工作表说明:")
    print(f"1. 完整性能矩阵 - 所有GPU×所有配置组合的性能矩阵")
    print(f"2-6. Seq4383~Seq70128 - 按序列长度分组的TP×DP性能表格")
    print(f"7. 最佳性能汇总 - 各GPU在不同序列长度下的最佳性能")
    print(f"8. TPDP组合统计 - 各TP×DP组合的统计分析")
    print(f"9. GPU最佳配置 - 每种GPU的最佳性能配置详情")
    print(f"10. 原始数据 - 完整的测试原始数据")
    print(f"11. 测试配置 - 测试参数和GPU列表")
    
    print("\n" + "=" * 80)
    print("🎉 Excel格式完整矩阵生成完成！")
    print("=" * 80)
    
    return results_df, performance_matrix, excel_filename

if __name__ == "__main__":
    results_df, performance_matrix, excel_file = main() 